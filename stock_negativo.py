"""
STOCK NEGATIVO — Ajuste automático en Tivendo (por sucursal)

Migrado desde el proyecto standalone tivendo_stock_negativo
(github.com/FoorKeM/tivendo-stock-negativo), agregando soporte multi-sucursal.

Flujo completo:
  1. PREPARAR (solo lectura):
     Login Tivendo → verificar empresa → Punto de Ventas →
     Informes > Informes de Inventario > Bodega de la sucursal activa >
     "Muestra artículos sin stock" + presentación "Resumido" > Exportar →
     Listado de artículos > Exportar →
     cruzar ambos informes (código A, stock negativo, código de barra
     presente, artículo activo y disponible para venta) → propuesta de ajuste.
  2. APLICAR (escribe en Tivendo):
     Inventario > Nuevo movimiento > AJUSTE POSITIVO DE INVENTARIO / ENTRADA /
     bodega de la sucursal → por cada artículo validado: buscar, revalidar
     stock actual contra la API, cargar la cantidad exacta para dejarlo en 0,
     Guardar con evidencia (JSON + captura).
  3. INFORME: Excel final solo con lo realmente ajustado.
  4. LIMPIEZA: borra los archivos temporales descargados.

Seguridad del flujo (heredada del proyecto original):
  - Solo procesa códigos que empiezan con "A".
  - Solo carga artículos cuyo stock actual es estrictamente negativo.
  - Valida código, nombre y código de barra antes de cargar cada artículo.
  - Vuelve a comprobar el stock real justo antes de cargarlo (por si cambió
    entre el paso de preparación y el de aplicación).
"""

import json
import re
import time
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from app_paths import AJUSTE_DIR, DESCARGA_DIR, runtime_path, configure_playwright_browsers
configure_playwright_browsers()
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.async_api import async_playwright

import config as _cfg
from utils import (
    asegurar_empresa_mercadohouse,
    asegurar_punto_ventas_abierto,
    crear_logger,
    crear_pagina_trabajo,
    guardar_diagnostico,
    login_tivendo,
    rotar_log,
)
from credenciales import TIVENDO_EMAIL as EMAIL, TIVENDO_PASSWORD as PASSWORD

# True = ver el navegador, False = correr invisible en segundo plano
MOSTRAR_NAVEGADOR = False

LOG_FILE = runtime_path("log_stock_negativo.txt")
log = crear_logger(LOG_FILE)

ULTIMA_PROPUESTA: list["AdjustmentItem"] | None = None
ULTIMO_RESULTADO: dict | None = None
ULTIMO_INFORME: str | None = None


@dataclass(frozen=True)
class AdjustmentItem:
    codigo: str
    descripcion: str
    stock_informe: str
    codigo_barra: str
    cantidad_propuesta: str


# ============================================================
#  Lectura y cruce de los Excel exportados (usa openpyxl: el análisis con
#  decimales y reparación de celdas numéricas es más sofisticado que el
#  parser manual del resto del proyecto)
# ============================================================

INVENTORY_HEADERS = ("CodArticulo", "Descripción Artículo", "Cantidad")
ARTICLE_HEADERS = (
    "Código",
    "Nombre",
    "Código barra interno",
    "Disponible para venta",
    "Activo",
)


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal(value) -> Decimal:
    if isinstance(value, Decimal):
        return value
    text = _text(value).replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Cantidad invalida: {value!r}") from exc


def _decimal_text(value: Decimal) -> str:
    return format(value, "f")


def _has_null_numbers(path: Path) -> bool:
    with zipfile.ZipFile(path, "r") as source:
        for name in source.namelist():
            if name.startswith("xl/worksheets/") and b"<v>null</v>" in source.read(name):
                return True
    return False


def _repair_null_numbers(path: Path) -> Path:
    handle = tempfile.NamedTemporaryFile(prefix="stock_reparado_", suffix=".xlsx", delete=False)
    repaired = Path(handle.name)
    handle.close()
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(repaired, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = data.replace(b"<v>null</v>", b"<v></v>")
            target.writestr(info, data)
    return repaired


def _read_rows(path: Path, required_headers: tuple[str, ...]):
    repaired_path = _repair_null_numbers(path) if _has_null_numbers(path) else None
    workbook = load_workbook(repaired_path or path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        for row in iterator:
            normalized = tuple(_text(value) for value in row)
            if all(header in normalized for header in required_headers):
                indexes = {header: normalized.index(header) for header in required_headers}
                return indexes, list(iterator)
    finally:
        workbook.close()
        if repaired_path:
            repaired_path.unlink(missing_ok=True)
    raise ValueError(
        f"No se encontraron las columnas requeridas en {path.name}: " + ", ".join(required_headers)
    )


def leer_negativos(path: Path) -> list[dict]:
    indexes, rows = _read_rows(path, INVENTORY_HEADERS)
    negativos = []
    for row in rows:
        code = _text(row[indexes["CodArticulo"]])
        if not code.upper().startswith("A"):
            continue
        quantity = _decimal(row[indexes["Cantidad"]])
        if quantity >= 0:
            continue
        negativos.append({
            "codigo": code,
            "descripcion": _text(row[indexes["Descripción Artículo"]]),
            "stock": quantity,
        })
    negativos.sort(key=lambda item: (item["stock"], item["codigo"]))
    return negativos


def leer_articulos(path: Path) -> dict[str, list[dict]]:
    indexes, rows = _read_rows(path, ARTICLE_HEADERS)
    articulos: dict[str, list[dict]] = {}
    for row in rows:
        code = _text(row[indexes["Código"]])
        if not code:
            continue
        articulos.setdefault(code, []).append({
            "nombre": _text(row[indexes["Nombre"]]),
            "codigo_barra": _text(row[indexes["Código barra interno"]]),
            "venta": _text(row[indexes["Disponible para venta"]]),
            "activo": _text(row[indexes["Activo"]]),
        })
    return articulos


def cruzar_negativos(
    inventory_path: Path, articles_path: Path
) -> tuple[list[AdjustmentItem], list[dict]]:
    """Cruza inventario negativo con el listado de artículos.

    Los artículos que no cumplen los requisitos (sin código de barra, sin
    coincidencia única, inactivos o no disponibles para venta) se omiten del
    ajuste en vez de abortar todo el proceso — se devuelven aparte para que
    se puedan revisar y corregir manualmente en Tivendo.
    """
    negativos = leer_negativos(inventory_path)
    articulos = leer_articulos(articles_path)
    omitidos = []
    resultado = []

    for negativo in negativos:
        coincidencias = articulos.get(negativo["codigo"], [])
        if len(coincidencias) != 1:
            omitidos.append({
                "codigo": negativo["codigo"],
                "descripcion": negativo["descripcion"],
                "motivo": f"se esperaban 1 coincidencia y hay {len(coincidencias)}",
            })
            continue
        articulo = coincidencias[0]
        if not articulo["codigo_barra"]:
            omitidos.append({
                "codigo": negativo["codigo"],
                "descripcion": negativo["descripcion"],
                "motivo": "sin Codigo barra interno",
            })
            continue
        if articulo["activo"].casefold() != "si":
            omitidos.append({
                "codigo": negativo["codigo"],
                "descripcion": negativo["descripcion"],
                "motivo": "articulo inactivo",
            })
            continue
        if articulo["venta"].casefold() != "si":
            omitidos.append({
                "codigo": negativo["codigo"],
                "descripcion": negativo["descripcion"],
                "motivo": "no disponible para venta",
            })
            continue

        cantidad = abs(negativo["stock"])
        resultado.append(AdjustmentItem(
            codigo=negativo["codigo"],
            descripcion=negativo["descripcion"],
            stock_informe=_decimal_text(negativo["stock"]),
            codigo_barra=articulo["codigo_barra"],
            cantidad_propuesta=_decimal_text(cantidad),
        ))

    return resultado, omitidos


def guardar_propuesta(
    inventory_path: Path,
    articles_path: Path,
    items: list[AdjustmentItem],
    bodega: str,
    omitidos: list[dict] | None = None,
) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destino = AJUSTE_DIR / f"propuesta_ajuste_{stamp}.json"
    payload = {
        "creada": datetime.now().isoformat(timespec="seconds"),
        "bodega": bodega,
        "inventario": str(inventory_path),
        "listado_articulos": str(articles_path),
        "cantidad_articulos": len(items),
        "articulos": [asdict(item) for item in items],
        "omitidos": omitidos or [],
    }
    destino.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return destino


# ============================================================
#  Exportar inventario y listado de artículos desde Tivendo POS
# ============================================================

async def _open_inventory_report(page) -> None:
    reports = page.get_by_text("Informes", exact=True)
    await reports.first.click()
    inventory = page.get_by_text("Informes de Inventario", exact=True)
    await inventory.wait_for(state="visible", timeout=15000)
    await inventory.click()
    await page.get_by_text("Bodegas", exact=True).wait_for(state="visible", timeout=30000)
    menu_toggle = page.locator("#menu_lateral button.boton_toggle_menu")
    if await menu_toggle.count():
        await menu_toggle.evaluate("element => element.click()")
        await page.wait_for_timeout(500)
    log("Informe de Inventario abierto.")


async def _select_single_warehouse(page) -> None:
    selector = page.locator('[role="combobox"][aria-labelledby*="selector_tipo_consulta_bodegas"]')
    await selector.wait_for(state="visible", timeout=10000)
    await selector.click()
    option = page.get_by_role("option", name=re.compile(r"Una en Particular", re.IGNORECASE))
    await option.wait_for(state="visible", timeout=10000)
    await option.click()


async def _choose_warehouse(page, bodega: str) -> None:
    await _select_single_warehouse(page)
    await page.wait_for_timeout(400)

    search = page.locator('input[type="text"]:not(.MuiSelect-nativeInput)').first
    await search.wait_for(state="visible", timeout=10000)
    await search.fill(bodega.split(" ")[0][:3].lower())
    option = page.get_by_text(bodega, exact=True)
    await option.wait_for(state="visible", timeout=15000)
    await option.click()
    log(f"Bodega seleccionada: {bodega}.")


async def _enable_out_of_stock(page) -> None:
    checkbox = page.locator('input[name="informes_de_ventas_muestra_articulos_sin_stock"]')
    await checkbox.wait_for(state="attached", timeout=10000)
    if not await checkbox.is_checked():
        await checkbox.check()
    log("Opcion 'Muestra articulos sin stock' activada.")


async def _ensure_summary(page) -> None:
    summary = page.locator(
        'input[name="radios_nivel_de_presentacion_informes_de_inventario"][value="Resumido"]'
    )
    await summary.wait_for(state="attached", timeout=10000)
    if not await summary.is_checked():
        await summary.check()
    log("Presentacion 'Resumido' confirmada.")


async def _exportar_inventario(page, bodega: str) -> Path:
    button = page.get_by_role("button", name="Exportar")
    await button.wait_for(state="visible", timeout=15000)
    log("Exportando informe de inventario...")
    async with page.expect_download(timeout=300000) as info:
        await button.click()
    download = await info.value
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = Path(download.suggested_filename or "inventario.xlsx").suffix or ".xlsx"
    destino = DESCARGA_DIR / f"Inventario_{bodega.replace(' ', '_')}_{stamp}{suffix}"
    await download.save_as(str(destino))
    if not destino.exists() or destino.stat().st_size == 0:
        raise RuntimeError("La descarga de inventario termino, pero el archivo esta vacio o no existe.")
    log(f"Informe de inventario descargado: {destino}")
    return destino


async def _exportar_listado_articulos(page) -> Path:
    log("Abriendo Listado de articulos...")
    await page.goto("https://tivendoapp.defontana.com/articulos/listado_articulos", timeout=45000)
    button = page.get_by_role("button", name="Exportar")
    await button.wait_for(state="visible", timeout=30000)
    log("Exportando listado de articulos...")
    async with page.expect_download(timeout=300000) as info:
        await button.click()
    download = await info.value
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    destino = DESCARGA_DIR / f"Listado_Articulos_{stamp}.xlsx"
    await download.save_as(str(destino))
    if not destino.exists() or destino.stat().st_size == 0:
        raise RuntimeError("El listado de articulos descargado esta vacio.")
    log(f"Listado de articulos descargado: {destino}")
    return destino


async def preparar():
    """Exporta inventario + listado de artículos y arma la propuesta de ajuste (solo lectura)."""
    suc = _cfg.sucursal_activa()
    bodega = suc["tivendo_bodega"]
    async with async_playwright() as p:
        browser, context, page = await crear_pagina_trabajo(
            p, mostrar_navegador=MOSTRAR_NAVEGADOR, carpeta_descarga=str(DESCARGA_DIR), incognito=True,
        )
        active_page = page
        try:
            await login_tivendo(page, EMAIL, PASSWORD, log)
            if "portal.defontana.com" in page.url:
                await asegurar_empresa_mercadohouse(page, log, EMAIL, PASSWORD)
                active_page = await asegurar_punto_ventas_abierto(context, page, log)

            await _open_inventory_report(active_page)
            await _choose_warehouse(active_page, bodega)
            await _enable_out_of_stock(active_page)
            await _ensure_summary(active_page)
            inventario_path = await _exportar_inventario(active_page, bodega)

            articulos_path = await _exportar_listado_articulos(active_page)
            items, omitidos = cruzar_negativos(inventario_path, articulos_path)
            propuesta_path = guardar_propuesta(inventario_path, articulos_path, items, bodega, omitidos)
            if omitidos:
                log(f"⚠ Articulos omitidos del ajuste ({len(omitidos)}) — no cumplen los requisitos:")
                for omitido in omitidos:
                    log(f"  {omitido['codigo']} | {omitido['descripcion']} | motivo: {omitido['motivo']}")
            log(f"Propuesta creada con {len(items)} articulo(s): {propuesta_path}")
            return inventario_path, articulos_path, propuesta_path, items, omitidos
        except Exception as exc:
            await guardar_diagnostico(active_page, exc, log, "stock_negativo_preparar")
            raise
        finally:
            await browser.close()


# ============================================================
#  Aplicar el ajuste (escribe en Tivendo)
# ============================================================

async def _select_option(page, selector: str, option: str) -> None:
    control = page.locator(selector)
    await control.wait_for(state="visible", timeout=30000)
    await control.click()
    target = page.get_by_role("option", name=option, exact=True)
    await target.wait_for(state="visible", timeout=15000)
    await target.click()


async def _open_movement(page) -> None:
    inventario = page.get_by_text("Inventario", exact=True).first
    await inventario.wait_for(state="visible", timeout=30000)
    await inventario.click()
    movimiento = page.get_by_text("Nuevo movimiento", exact=True)
    await movimiento.wait_for(state="visible", timeout=15000)
    await movimiento.click()
    await page.locator("#selector_selector_documento_inventario").wait_for(state="visible", timeout=30000)
    toggle = page.locator("#menu_lateral button.boton_toggle_menu")
    if await toggle.count():
        await toggle.evaluate("element => element.click()")
        await page.wait_for_timeout(500)


async def _configure_movement(page, bodega: str) -> None:
    await _select_option(page, "#selector_selector_documento_inventario", "AJUSTE POSITIVO DE INVENTARIO")
    await _select_option(page, "#selector_selector_motivo_movimiento_inventario", "ENTRADA")
    await _select_option(page, '[id*="inventario_bodega_"][role="combobox"]', bodega)
    textarea = page.locator("textarea").first
    await textarea.fill(f"Regularizacion stock negativo {bodega} - {datetime.now():%Y-%m-%d}")


async def _wait_catalog(page) -> None:
    search = page.locator('input[placeholder*="digo o nombre"]')
    await search.wait_for(state="visible", timeout=30000)
    indicator = page.get_by_role("button", name="Tivendo SOS", exact=False)
    if await indicator.count():
        started = time.monotonic()
        log("Esperando que Tivendo prepare el catalogo de articulos...")
        try:
            next_message = 10
            while True:
                optimizing = await page.evaluate(
                    """
                    () => {
                        const button = [...document.querySelectorAll('button')]
                            .find(el => (el.getAttribute('aria-label') || '').includes('Tivendo SOS'));
                        if (!button) return false;
                        return /Optimizando/i.test(button.getAttribute('aria-label') || '');
                    }
                    """
                )
                if not optimizing:
                    break
                elapsed = int(time.monotonic() - started)
                if elapsed >= 180:
                    raise TimeoutError
                if elapsed >= next_message:
                    log(f"Tivendo sigue preparando el catalogo ({elapsed} s)...")
                    next_message += 10
                await page.wait_for_timeout(2000)
        except Exception:
            log("Tivendo sigue optimizando; se intentara usar el buscador.")
        else:
            elapsed = int(time.monotonic() - started)
            log(f"Catalogo de Tivendo listo ({elapsed} s).")


def _parse_stock(value: str) -> Decimal:
    normalizado = value.strip().replace(",", ".")
    try:
        return Decimal(normalizado)
    except InvalidOperation as exc:
        raise RuntimeError(f"No se pudo leer el stock actual: {value!r}") from exc


async def _search_product(page, item: AdjustmentItem, zona_stock: str):
    fila_vacia = page.locator("tr.articulos_inventario_tr").filter(
        has=page.locator("td.sin_seleccionar_articulo")
    ).last
    buscador = fila_vacia.locator('input[placeholder*="digo o nombre"]')
    await buscador.wait_for(state="visible", timeout=30000)

    async with page.expect_response(
        lambda response: "/producto/paginado" in response.url, timeout=120000
    ) as response_info:
        await buscador.fill(item.codigo)
    response = await response_info.value
    payload = await response.json()
    productos = payload.get("productos") or []
    coincidencias = [p for p in productos if p.get("sku") == item.codigo]
    if len(coincidencias) != 1:
        raise RuntimeError(f"Tivendo devolvio {len(coincidencias)} coincidencias para {item.codigo}.")

    producto = coincidencias[0]
    if str(producto.get("codigoInterno") or "") != item.codigo_barra:
        raise RuntimeError(
            f"El codigo de barra de {item.codigo} cambio en Tivendo: {producto.get('codigoInterno')!r}."
        )
    if str(producto.get("nombre") or "").strip() != item.descripcion:
        raise RuntimeError(f"El nombre de {item.codigo} no coincide con el informe.")

    stock_data = producto.get("stock") or {}
    bodegas_con_stock = stock_data.get("bodegasConStock") or []
    stock_bodega = None
    for bodega_stock in bodegas_con_stock:
        if bodega_stock.get("idBodega") == zona_stock:
            stock_bodega = bodega_stock.get("stock")
            break
    if stock_bodega is None:
        zonas_encontradas = sorted({str(b.get("idBodega")) for b in bodegas_con_stock})
        raise RuntimeError(
            f"Tivendo no devolvio stock de la zona '{zona_stock}' para {item.codigo}. "
            f"Zonas encontradas en la respuesta: {zonas_encontradas or 'ninguna'}. "
            "Corrige 'tivendo_zona_stock' en config.py para esta sucursal con el valor real."
        )
    stock = _parse_stock(str(stock_bodega))
    if stock >= 0:
        await buscador.fill("")
        return None, stock

    codigo_loc = page.locator(f'li[role="menuitem"] h6.codigo_articulo[aria-label="{item.codigo}"]')
    await codigo_loc.wait_for(state="visible", timeout=15000)
    cantidades_habilitadas = page.locator('input[type="number"]:not([disabled])')
    conteo_previo = await cantidades_habilitadas.count()
    await codigo_loc.click()
    await page.wait_for_function(
        """expected => document.querySelectorAll('input[type="number"]:not([disabled])').length > expected""",
        arg=conteo_previo,
        timeout=15000,
    )

    filas_seleccionadas = page.locator("tr.articulos_inventario_tr").filter(
        has=page.locator('input[type="number"]:not([disabled])')
    )
    await filas_seleccionadas.first.wait_for(state="visible", timeout=15000)
    placeholders = []
    for indice in range(await filas_seleccionadas.count()):
        candidata = filas_seleccionadas.nth(indice)
        campo_producto = candidata.locator('input[type="text"]').first
        placeholder = await campo_producto.get_attribute("placeholder") or ""
        placeholders.append(placeholder)
        if placeholder == item.descripcion:
            return candidata, stock
    raise RuntimeError(
        f"No se encontro la fila seleccionada de {item.codigo} despues de agregarla. "
        f"Filas visibles: {placeholders!r}"
    )


async def _load_item(page, item: AdjustmentItem, zona_stock: str) -> dict:
    fila, stock_api = await _search_product(page, item, zona_stock)
    if fila is None:
        return {
            "codigo": item.codigo,
            "codigo_barra": item.codigo_barra,
            "stock_actual": format(stock_api, "f"),
            "estado": "omitido_no_negativo",
        }

    stock_display = fila.locator(".informacion_stock_articulo")
    stock_text = await stock_display.get_attribute("aria-label") or ""
    stock_mostrado = _parse_stock(stock_text)
    if stock_mostrado != stock_api:
        raise RuntimeError(
            f"El stock mostrado de {item.codigo} ({stock_mostrado}) no coincide con la API ({stock_api})."
        )
    if stock_mostrado >= 0:
        raise RuntimeError(f"{item.codigo} dejo de estar negativo durante la seleccion.")

    campo_cantidad = fila.locator('input[type="number"]:not([disabled])')
    cantidad = format(abs(stock_mostrado), "f")
    await campo_cantidad.fill(cantidad)
    return {
        "codigo": item.codigo,
        "codigo_barra": item.codigo_barra,
        "stock_actual": stock_text,
        "cantidad": cantidad,
        "estado": "cargado",
    }


def _guardar_resultado(payload: dict) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = AJUSTE_DIR / f"resultado_ajuste_{stamp}.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


async def aplicar(items: list[AdjustmentItem], simular: bool = False) -> Path:
    """Aplica el ajuste en Tivendo para los artículos ya validados por preparar()."""
    suc = _cfg.sucursal_activa()
    bodega = suc["tivendo_bodega"]
    zona_stock = suc["tivendo_zona_stock"]

    async with async_playwright() as p:
        browser, context, page = await crear_pagina_trabajo(p, mostrar_navegador=MOSTRAR_NAVEGADOR, incognito=True)
        active_page = page
        resultados = []
        try:
            await login_tivendo(page, EMAIL, PASSWORD, log)
            if "portal.defontana.com" in page.url:
                await asegurar_empresa_mercadohouse(page, log, EMAIL, PASSWORD)
                active_page = await asegurar_punto_ventas_abierto(context, page, log)

            log("Abriendo Nuevo movimiento...")
            await _open_movement(active_page)
            log("Configurando ajuste positivo, entrada y bodega...")
            await _configure_movement(active_page, bodega)
            log("Movimiento configurado.")
            await _wait_catalog(active_page)

            for indice, item in enumerate(items, start=1):
                log(f"Validando articulo {indice}/{len(items)}: {item.codigo}")
                resultados.append(await _load_item(active_page, item, zona_stock))

            cargados = [r for r in resultados if r["estado"] == "cargado"]
            if not cargados:
                raise RuntimeError("No quedan articulos con stock actual negativo.")

            payload = {
                "fecha": datetime.now().isoformat(timespec="seconds"),
                "simulacion": simular,
                "bodega": bodega,
                "articulos": resultados,
            }
            if simular:
                payload["estado"] = "simulado_sin_guardar"
                await active_page.screenshot(path=str(AJUSTE_DIR / "ultima_simulacion.png"), full_page=True)
                return _guardar_resultado(payload)

            payload["estado"] = "guardado_pendiente_confirmacion"
            payload["guardar_iniciado"] = datetime.now().isoformat(timespec="seconds")
            result_path = _guardar_resultado(payload)
            boton = active_page.get_by_role("button", name="Guardar", exact=True)
            respuestas = []

            def registrar_respuesta(response):
                if response.request.method in {"POST", "PUT"} and "api-pos-prod.defontana.com" in response.url:
                    respuestas.append({"url": response.url, "status": response.status, "ok": response.ok})

            active_page.on("response", registrar_respuesta)
            try:
                await boton.click()
                await active_page.wait_for_function(
                    """() => document.querySelectorAll('input[type="number"]:not([disabled])').length === 0""",
                    timeout=120000,
                )
            except Exception:
                payload["estado"] = "guardado_incierto_no_reintentar"
                payload["respuestas_guardado"] = respuestas
                result_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
                raise

            await active_page.wait_for_timeout(1500)
            payload["estado"] = "guardado"
            payload["confirmacion_guardado"] = "filas_cargadas_limpiadas"
            payload["respuestas_guardado"] = respuestas
            payload["url_final"] = active_page.url
            payload["texto_confirmacion"] = (await active_page.locator("body").inner_text())[-2000:]
            await active_page.screenshot(path=str(AJUSTE_DIR / "ultimo_ajuste_guardado.png"), full_page=True)
            result_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            return result_path
        except Exception as exc:
            await guardar_diagnostico(active_page, exc, log, "stock_negativo_aplicar")
            raise
        finally:
            await browser.close()


def limpiar_descargas(result_path: Path) -> list[str]:
    borrados = []
    errores = []
    for path in DESCARGA_DIR.iterdir():
        if not path.is_file():
            continue
        try:
            path.unlink()
            borrados.append(str(path))
        except OSError as exc:
            errores.append(f"{path}: {exc}")

    payload = json.loads(result_path.read_text(encoding="utf-8"))
    payload["limpieza_descargas"] = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "archivos_eliminados": borrados,
        "errores": errores,
    }
    result_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    if errores:
        raise RuntimeError(
            "El ajuste se completo, pero no se pudieron eliminar todas las descargas:\n- "
            + "\n- ".join(errores)
        )
    return borrados


# ============================================================
#  Informe Excel final
# ============================================================

INFORME_HEADERS = (
    "Fecha ajuste",
    "Bodega",
    "CodArticulo",
    "Descripcion Articulo",
    "Codigo barra interno",
    "Stock negativo",
    "Cantidad ajustada",
)


def _numero(value: str):
    decimal = Decimal(str(value).strip().replace(",", "."))
    if decimal == decimal.to_integral():
        return int(decimal)
    return float(decimal)


def _nombre_informe_por_defecto() -> str:
    return f"Informe_Stock_Negativo_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def crear_informe(result_path: Path, items: list[AdjustmentItem], destino: Path) -> Path:
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    descripciones = {item.codigo: item.descripcion for item in items}
    filas = []

    for articulo in payload.get("articulos", []):
        if articulo.get("estado") != "cargado":
            continue
        stock = Decimal(str(articulo["stock_actual"]).strip().replace(",", "."))
        cantidad = Decimal(str(articulo["cantidad"]).strip().replace(",", "."))
        if stock >= 0 or cantidad <= 0:
            continue
        codigo = str(articulo["codigo"]).strip()
        if not codigo.upper().startswith("A"):
            continue
        filas.append([
            payload["fecha"],
            payload["bodega"],
            codigo,
            descripciones.get(codigo, ""),
            str(articulo["codigo_barra"]).strip(),
            _numero(str(stock)),
            _numero(str(cantidad)),
        ])

    if not filas:
        raise RuntimeError("No hay articulos negativos ajustados para el informe.")

    destino = destino.with_suffix(".xlsx")
    destino.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Stock negativo"
    hoja.append(INFORME_HEADERS)
    for fila in filas:
        hoja.append(fila)

    encabezado = hoja[1]
    for celda in encabezado:
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for celda in hoja["E"][1:]:
        celda.number_format = "@"
    for fila in hoja.iter_rows(min_row=2, min_col=6, max_col=7):
        for celda in fila:
            celda.number_format = "0.########"

    anchos = {"A": 22, "B": 20, "C": 16, "D": 55, "E": 24, "F": 16, "G": 18}
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    hoja.freeze_panes = "A2"
    tabla = Table(displayName="AjustesStockNegativo", ref=f"A1:G{hoja.max_row}")
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    hoja.add_table(tabla)
    workbook.save(destino)
    workbook.close()

    verificacion = load_workbook(destino, read_only=True, data_only=True)
    try:
        if verificacion.active.max_row != len(filas) + 1:
            raise RuntimeError("El informe Excel no contiene todas las filas esperadas.")
    finally:
        verificacion.close()
    return destino


def elegir_destino_informe() -> Path | None:
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.wm_attributes("-topmost", True)
    try:
        seleccionado = filedialog.asksaveasfilename(
            parent=root,
            title="Guardar informe de stock negativo",
            defaultextension=".xlsx",
            initialfile=_nombre_informe_por_defecto(),
            filetypes=[("Libro de Excel", "*.xlsx")],
        )
    finally:
        root.destroy()
    return Path(seleccionado) if seleccionado else None


def guardar_informe_interactivo(result_path: Path, items: list[AdjustmentItem]) -> Path:
    destino = elegir_destino_informe()
    if destino is None:
        destino = AJUSTE_DIR / _nombre_informe_por_defecto()
        log(f"No se eligio una carpeta para el informe. Se guardara en: {destino}")
    return crear_informe(result_path, items, destino)


# ============================================================
#  Orquestador de alto nivel para menu.py
# ============================================================

async def ejecutar_ajuste_completo(simular: bool = False, confirmar_fn=None) -> bool:
    """Prepara, muestra la propuesta, aplica (o simula) y genera el informe.

    confirmar_fn(items, omitidos) -> bool, si se entrega, se llama despues de preparar()
    y antes de aplicar() de verdad; si devuelve False se aborta sin tocar Tivendo.
    """
    global ULTIMA_PROPUESTA, ULTIMO_RESULTADO, ULTIMO_INFORME
    ULTIMA_PROPUESTA = None
    ULTIMO_RESULTADO = None
    ULTIMO_INFORME = None
    rotar_log(LOG_FILE)

    log("=" * 50)
    log("INICIO AJUSTE DE STOCK NEGATIVO")
    suc = _cfg.sucursal_activa()
    log(f"🏬 Sucursal activa : {suc['nombre']}")
    log(f"   Bodega Tivendo  : {suc['tivendo_bodega']}")
    log("=" * 50)

    inventario_path, articulos_path, propuesta_path, items, omitidos = await preparar()
    ULTIMA_PROPUESTA = items

    if omitidos:
        log(f"⚠ Articulos con stock negativo OMITIDOS del ajuste ({len(omitidos)}):")
        for omitido in omitidos:
            log(f"  {omitido['codigo']} | {omitido['descripcion']} | motivo: {omitido['motivo']}")
        log("  (revisa estos manualmente en Tivendo; no se les hizo ningun cambio)")

    if not items:
        log("✓ No hay articulos con stock negativo para ajustar (revisa los omitidos arriba si los hay).")
        return True

    log(f"Articulos con stock negativo a ajustar: {len(items)}")
    for item in items:
        log(f"  {item.codigo} | {item.descripcion} | stock: {item.stock_informe} | ajuste: +{item.cantidad_propuesta}")

    if confirmar_fn is not None and not confirmar_fn(items, omitidos):
        log("Ajuste cancelado por el usuario antes de aplicarlo. No se modifico Tivendo.")
        return False

    result_path = await aplicar(items, simular=simular)
    ULTIMO_RESULTADO = json.loads(result_path.read_text(encoding="utf-8"))

    if simular:
        log("SIMULACION COMPLETADA — no se modifico stock.")
        return True

    log("✓ Ajuste guardado en Tivendo.")
    informe_path = guardar_informe_interactivo(result_path, items)
    ULTIMO_INFORME = str(informe_path)
    log(f"Informe guardado: {informe_path}")

    eliminados = limpiar_descargas(result_path)
    log(f"Descargas eliminadas: {len(eliminados)}")
    return True
